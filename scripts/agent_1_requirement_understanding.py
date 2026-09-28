import os
import sys
import csv
import re
import json
import subprocess
import requests
from html import unescape
from typing import List, Dict, Any

# Disable SSL Warnings for corporate environments (verify=False)
import urllib3
urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)

# Config/Constants
# ==============================================================================
# SECURITY & ACCESS CONSTRAINT:
# This agent operates strictly under a READ-ONLY model. Under no circumstances 
# does it execute write, create, or update commands back to DOORS Next Gen (DNG).
# All DNG interaction is restricted to query-based reading.
# ==============================================================================
DNG_SERVER_RAW = os.getenv("DNG_SERVER", "https://rb-alm-20-p.de.bosch.com")
# Strip trailing slashes and remove duplicate '/rm' or '/rm/' suffix
DNG_SERVER = DNG_SERVER_RAW.rstrip("/")
if DNG_SERVER.endswith("/rm"):
    DNG_SERVER = DNG_SERVER[:-3]
elif DNG_SERVER.endswith("rm"):
    # handles cases without slash
    DNG_SERVER = DNG_SERVER[:-2]
DNG_SERVER = DNG_SERVER.rstrip("/")

DNG_USERNAME = os.getenv("DNG_USERNAME")
DNG_PASSWORD = os.getenv("DNG_PASSWORD")

import config

OSLCQUERY_PATH = config.OSLCQUERY_PATH
TEMP_DIR = config.WORKSPACE_DIR

def clean_html(raw_html: str) -> str:
    if not raw_html:
        return ""
    text = unescape(raw_html)
    text = re.sub('<[^<]+?>', '', text)
    text = re.sub(r'\s+', ' ', text)
    return text.strip()

def cleanup_stale_cookies() -> None:
    """Deletes stale .cookies files in CWD or script directory to prevent elmclient 401 login bugs."""
    cookie_paths = [
        os.path.join(os.getcwd(), ".cookies"),
        os.path.join(os.path.dirname(os.path.abspath(__file__)), ".cookies"),
        r"D:\Hackathon\pmt-server\pmt-mcp-servers\mcp\dng\.cookies"
    ]
    for path in cookie_paths:
        if os.path.exists(path):
            try:
                os.remove(path)
                print(f"[+] Cleaned up stale session cookies at: {path}")
            except Exception as e:
                print(f"[-] Warning: Could not remove stale cookies file {path}: {e}")

def fetch_module_from_dng(project: str, component: str, configuration: str, module_name: str, output_csv: str) -> bool:
    """Runs oslcquery.EXE to download all artifacts belonging to the module."""
    if not os.path.isfile(OSLCQUERY_PATH):
        print(f"[-] Error: oslcquery.EXE not found at {OSLCQUERY_PATH}")
        return False
    
    if not DNG_USERNAME or not DNG_PASSWORD:
        print("[-] Error: DNG_USERNAME and DNG_PASSWORD environment variables must be set.")
        return False

    # Force cleanup of stale cookies to avoid 401 error crashes
    cleanup_stale_cookies()

    env = os.environ.copy()
    env["QUERY_USER"] = DNG_USERNAME
    env["QUERY_PASSWORD"] = DNG_PASSWORD

    cmd = [
        OSLCQUERY_PATH,
        "-J", DNG_SERVER,
        "-p", project,
        "-V",
        "-C", component,
        "-F", configuration,
        "-q", f'rm:module=^"{module_name}"',
        "-s", "*",
        "-O", output_csv
    ]

    print(f"[*] Querying DOORS Next Gen for module '{module_name}' in project '{project}'...")
    print(f"[*] Running: {' '.join(cmd)}")
    
    try:
        res = subprocess.run(cmd, env=env, capture_output=True, text=True, timeout=600)
        if res.returncode == 0:
            print("[+] Successfully fetched module requirements from DNG.")
            return True
        else:
            print(f"[-] DNG Fetch failed. Exit Code: {res.returncode}")
            print(f"[-] STDERR: {res.stderr}")
            return False
    except subprocess.TimeoutExpired:
        print("[-] DNG Fetch timed out after 10 minutes.")
        return False

def read_rows_from_file(file_path: str) -> List[Dict[str, Any]]:
    """Reads rows from either a CSV or XLSX file and returns them as a list of dictionaries."""
    if not os.path.exists(file_path):
        print(f"[-] File not found: {file_path}")
        return []

    if file_path.lower().endswith(".xlsx"):
        print(f"[*] Reading Excel file using openpyxl: {file_path}")
        try:
            import openpyxl
            wb = openpyxl.load_workbook(file_path, data_only=True)
            sheet = wb.active
            rows = []
            headers = []
            for r_idx, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                if r_idx == 1:
                    headers = [str(cell).strip() if cell is not None else f"Col{idx}" for idx, cell in enumerate(row, start=1)]
                    continue
                if all(cell is None for cell in row):
                    continue
                row_dict = {}
                for idx, cell in enumerate(row):
                    val = str(cell) if cell is not None else ""
                    header_name = headers[idx] if idx < len(headers) else f"Col{idx+1}"
                    row_dict[header_name] = val
                rows.append(row_dict)
            return rows
        except Exception as e:
            print(f"[-] Failed to read XLSX file {file_path}: {e}")
            return []
    else:
        # Default to CSV behavior
        print(f"[*] Reading CSV file: {file_path}")
        rows = []
        # Attempt UTF-8, fallback to cp1252 or utf-8-sig
        try:
            with open(file_path, "r", encoding="utf-8-sig") as f:
                reader = csv.DictReader(f)
                for row in reader:
                    rows.append(row)
        except UnicodeDecodeError:
            try:
                with open(file_path, "r", encoding="cp1252") as f:
                    reader = csv.DictReader(f)
                    for row in reader:
                        rows.append(row)
            except Exception as e:
                print(f"[-] Failed to read CSV file {file_path}: {e}")
                return []
        return rows

def extract_section_requirements(file_path: str, target_section: str) -> List[Dict[str, Any]]:
    """Reads the module file (CSV or XLSX), finds the target heading, and extracts all child requirements."""
    print(f"[*] Extracting requirements for section '{target_section}' from file...")
    
    rows = read_rows_from_file(file_path)
    if not rows:
        return []

    section_started = False
    extracted_reqs = []
    target_section_num = None

    for idx, row in enumerate(rows):
        # Flexible mapping to support both oslcquery CSV format and DOORS-downloaded Excel format
        ident = (row.get("Identifier") or row.get("id") or "").strip()
        shape = (row.get("oslc:instanceShape") or row.get("Artifact Type") or "").strip()
        
        is_heading_col = row.get("isHeading", "")
        is_heading = (is_heading_col in [True, "True", "true"]) or (shape == "Heading")
        
        primary_text = clean_html(row.get("Primary Text", ""))
        
        # Heading titles in DOORS UI exports typically reside in the "Primary Text" field,
        # while they may be in the "Title" field in standard OSLC CSV outputs.
        title = (row.get("Title") or "").strip()
        if not title:
            if is_heading:
                title = primary_text
            else:
                # Use first line/sentence of primary_text as a fallback title for requirements
                first_line = primary_text.split('\n')[0].strip()
                if len(first_line) > 80:
                    title = first_line[:77] + "..."
                else:
                    title = first_line

        section_num = str(row.get("section") or "").strip()
        
        # Check if we hit our target heading
        if is_heading and target_section.lower() in title.lower():
            section_started = True
            target_section_num = section_num
            print(f"[+] Found target heading: '{title}' (ID {ident}) at row {idx+2} (Section: {target_section_num})")
            continue

        if section_started:
            # Stop extraction when we hit a sibling or parent heading
            if is_heading:
                # 1. Hierarchy/section-number based stopping condition (Highly precise)
                if target_section_num and section_num:
                    prefix = target_section_num + "."
                    if section_num != target_section_num and not section_num.startswith(prefix):
                        print(f"[*] Encountered sibling/parent section heading '{title}' (Section: {section_num}). Stopping extraction.")
                        break
                else:
                    # 2. Hardcoded fallback major sibling headings (Only used when section numbers are not available)
                    major_sibling_headings = [
                        "swupdate", "query interface", "bank sync", "ota", "main domain sail updater"
                    ]
                    if title.lower() in major_sibling_headings:
                        print(f"[*] Encountered sibling major section heading '{title}' (ID {ident}). Stopping extraction.")
                        break
            
            # We want to parse Software Requirements (or information blocks containing logic)
            if shape in ["Software Requirement", "Software Req", "Requirement", "Information"]:
                extracted_reqs.append({
                    "id": ident,
                    "title": title,
                    "type": shape,
                    "content": primary_text,
                    "comment": row.get("RB_Comment_BoschInternal", "")
                })

    print(f"[+] Extracted {len(extracted_reqs)} requirements under section '{target_section}'.")
    return extracted_reqs

def run_gemini_cli(cmd: list, timeout: int = 300) -> str:
    """Invokes gemini.cmd, streams output to console in real-time to prevent timeouts, and returns full output."""
    try:
        process = subprocess.Popen(
            cmd,
            stdout=subprocess.PIPE,
            stderr=subprocess.STDOUT,
            text=True,
            encoding="utf-8",
            bufsize=1
        )
        
        full_output = []
        for line in iter(process.stdout.readline, ""):
            # Print to terminal live to keep the connection active and let the user see progress
            print(line, end="")
            full_output.append(line)
            
        process.stdout.close()
        process.wait(timeout=timeout)
        return "".join(full_output)
    except Exception as e:
        print(f"[-] Gemini CLI execution failed: {e}")
        return ""

def decode_with_llm(requirements: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    """Sends the requirements to local gemini CLI to extract functional behaviors and testcases."""
    decoded_results = []
    
    # Process requirements in small batches to fit context and limit payload size
    batch_size = 18
    for batch_idx in range(0, len(requirements), batch_size):
        batch = requirements[batch_idx : batch_idx + batch_size]
        print(f"[*] LLM Decoding Batch {batch_idx//batch_size + 1} of {(len(requirements)-1)//batch_size + 1}...")

        prompt = f"""
You are the "Requirement Understanding Agent" (Agent 1) in a safety-critical automotive systems test-generation pipeline.
Your task is to analyze the following Software Requirements, extract their functional behaviors, and determine all possible test scenarios (positive, negative, and boundary) that can be generated to verify them.

Analyze each requirement and generate:
- Positive Scenarios (successful execution paths)
- Negative Scenarios (erroneous behavior, invalid inputs, failure conditions)
- Boundary Scenarios (limits, rollback conditions, timeouts, retry limits)

Here are the requirements to analyze:
{json.dumps(batch, indent=2)}

You MUST output your response STRICTLY as a valid JSON array of objects conforming to this schema:
[
  {{
    "requirement_id": "string (the DNG ID, e.g. 2469074)",
    "title": "string (the requirement title)",
    "functional_scope": "string (e.g. Bank Swap, Integrity Validation)",
    "scenarios": [
      {{
        "scenario_type": "string ('positive', 'negative', or 'boundary')",
        "description": "string (short description of the test objective)",
        "preconditions": ["string (state of system before trigger)"],
        "action": "string (the event/command/trigger to test)",
        "expected_outcome": "string (expected response/state)"
      }}
    ]
  }}
]

Output ONLY raw valid JSON. Do not wrap in backticks or markdown formatting.
"""

        try:
            # Write the prompt to a temporary file to avoid command-line and stdin encoding/forwarding issues on Windows
            temp_prompt_path = os.path.join(TEMP_DIR, f"temp_prompt_batch_{batch_idx}.txt")
            with open(temp_prompt_path, "w", encoding="utf-8") as pf:
                pf.write(prompt)

            # Call gemini CLI natively
            cmd = [
                config.GEMINI_CMD,
                "-p", f"Analyze the requirements in @{temp_prompt_path} and output strictly raw valid JSON array conforming to the requested schema contract."
            ]
            
            llm_response = run_gemini_cli(cmd)
            
            # Clean up the temp prompt file
            try:
                os.remove(temp_prompt_path)
            except Exception:
                pass
            
            # Extract JSON from response
            llm_response = llm_response.strip()
            
            # Find the starting [ or { and ending ] or } to isolate JSON from logs/stdout clutter
            json_match = re.search(r'(\[.*\]|\{.*\})', llm_response, re.DOTALL)
            if json_match:
                json_str = json_match.group(1)
            else:
                json_str = llm_response

            batch_decoded = json.loads(json_str)
            if isinstance(batch_decoded, dict) and "requirements" in batch_decoded:
                batch_decoded = batch_decoded["requirements"]
            if isinstance(batch_decoded, list):
                decoded_results.extend(batch_decoded)
            else:
                decoded_results.append(batch_decoded)

        except Exception as e:
            print(f"[-] Error calling LLM for batch {batch_idx//batch_size + 1}: {e}")
            # Fallback with placeholders
            for req in batch:
                decoded_results.append({
                    "requirement_id": req["id"],
                    "title": req["title"],
                    "functional_scope": "Unknown",
                    "scenarios": []
                })

    return decoded_results

def main():
    import argparse
    parser = argparse.ArgumentParser(description="Agent 1: Requirement Understanding Agent")
    parser.add_argument("--project", default="NEO", help="DNG Project Area Name")
    parser.add_argument("--component", default="NEO", help="DNG Component Name")
    parser.add_argument("--config", default="DNG NEO Stream", help="DNG Configuration/Stream Name")
    parser.add_argument("--module", default=None, help="DNG Module Title (omitted for interactive prompt)")
    parser.add_argument("--section", default=None, help="The section heading title to extract (omitted for interactive prompt)")
    parser.add_argument("--offline", action="store_true", help="Bypass DNG connection and use local CSV/XLSX if exists")
    parser.add_argument("--input-file", default=None, help="Path to local Excel (.xlsx) or CSV file downloaded from DNG (used for offline mode)")
    args = parser.parse_args()

    is_offline = args.offline or (args.input_file is not None)

    if is_offline and args.input_file is None:
        parser.error("--input-file is mandatory when running in offline mode.")

    # Interactive User Prompts for Module & Section (with safe defaults)
    module_title = args.module
    if not is_offline and module_title is None:
        try:
            user_input = input("Enter DNG Module Title (default: 'SWRS Update'): ").strip()
            module_title = user_input if user_input else "SWRS Update"
        except (KeyboardInterrupt, EOFError):
            print("\n[-] Cancelled by user.")
            sys.exit(1)

    section_title = args.section
    if section_title is None:
        try:
            user_input = input("Enter Section Heading Title (default: 'MCU'): ").strip()
            section_title = user_input if user_input else "MCU"
        except (KeyboardInterrupt, EOFError):
            print("\n[-] Cancelled by user.")
            sys.exit(1)

    if is_offline:
        input_file = args.input_file
    else:
        input_file = args.input_file if args.input_file is not None else os.path.join(TEMP_DIR, "swrs_update_requirements.csv")

    output_json_path = os.path.join(TEMP_DIR, "decoded_requirements.json")

    # Step 1: Connect to DNG and Fetch module requirements to CSV (unless offline mode is selected)
    if not is_offline:
        success = fetch_module_from_dng(args.project, args.component, args.config, module_title, input_file)
        if not success:
            print("[-] DNG Fetch failed. Aborting.")
            sys.exit(1)
    else:
        if not os.path.exists(input_file):
            print(f"[-] Error: Specified input file '{input_file}' not found.")
            sys.exit(1)

    # Step 2: Parse and Extract the target section requirements
    extracted_reqs = extract_section_requirements(input_file, section_title)
    if not extracted_reqs:
        print(f"[-] No requirements extracted under section '{section_title}'. Aborting.")
        sys.exit(1)

    # Step 3: Feed requirements into LLM for decoding positive, negative, and boundary scenarios
    decoded_reqs = decode_with_llm(extracted_reqs)

    # Step 4: Save structured output
    with open(output_json_path, "w", encoding="utf-8") as out:
        json.dump(decoded_reqs, out, indent=2)

    print(f"\n[+] SUCCESS: Agent 1 has decoded {len(decoded_reqs)} requirements!")
    print(f"[+] Output JSON saved to: {output_json_path}")
    print("[+] Ready for ingestion by Agent 2 (Mapping) and Agent 3 (Gap Analyzer)!")

if __name__ == "__main__":
    main()
