import os
import sys
import openpyxl
import config

if len(sys.argv) < 2:
    print(f"[-] Error: Missing required file path argument.")
    print(f"[*] Usage: python {os.path.basename(sys.argv[0])} <path_to_excel_file.xlsx>")
    sys.exit(1)

path = sys.argv[1]

if not os.path.exists(path):
    print(f"[-] Error: File not found: {path}")
    sys.exit(1)
wb = openpyxl.load_workbook(path, data_only=True)
sheet = wb.active

print(f"Sheet Name: {sheet.title}")
print(f"Max Row: {sheet.max_row}")
print(f"Max Col: {sheet.max_column}")

headers = [str(cell).strip() if cell is not None else f"Col{idx}" for idx, cell in enumerate(next(sheet.iter_rows(values_only=True)), start=1)]
print(f"Headers: {headers}")

heading_rows = []
for r_idx, row in enumerate(sheet.iter_rows(values_only=True), start=1):
    if r_idx == 1:
        continue
    shape = str(row[headers.index("Artifact Type")] if "Artifact Type" in headers else "")
    title = str(row[headers.index("Primary Text")] if "Primary Text" in headers else "")
    ident = str(row[headers.index("id")] if "id" in headers else "")
    
    if "heading" in shape.lower() or shape == "Heading":
        heading_rows.append((r_idx, ident, title))

print(f"\nTotal Headings Found: {len(heading_rows)}")
print("First 15 Headings:")
for idx, h in enumerate(heading_rows[:15]):
    print(f"  - Row {h[0]}: ID {h[1]} - '{h[2]}'")
if len(heading_rows) > 15:
    print(f"  - ... and {len(heading_rows) - 15} more headings")

# Now let's trace the "MCU" heading and count how many requirements are under it
mcu_index = -1
for i, h in enumerate(heading_rows):
    if "mcu" in h[2].lower():
        mcu_index = i
        print(f"\n[+] Target MCU Heading found at heading index {i}: Row {h[0]}, ID {h[1]} - '{h[2]}'")
        break

if mcu_index != -1:
    start_row = heading_rows[mcu_index][0]
    end_row = sheet.max_row
    if mcu_index + 1 < len(heading_rows):
        end_row = heading_rows[mcu_index + 1][0] - 1
        
    print(f"MCU Section range: Rows {start_row} to {end_row}")
    
    reqs_in_mcu = 0
    types_in_mcu = {}
    for r_idx in range(start_row + 1, end_row + 1):
        row_cells = [cell.value for cell in sheet[r_idx]]
        shape = str(row_cells[headers.index("Artifact Type")] if "Artifact Type" in headers else "")
        if shape in ["Software Requirement", "Software Req", "Requirement", "Information"]:
            reqs_in_mcu += 1
            types_in_mcu[shape] = types_in_mcu.get(shape, 0) + 1
            
    print(f"Number of requirements under MCU section: {reqs_in_mcu}")
    print(f"Requirement types distribution: {types_in_mcu}")
