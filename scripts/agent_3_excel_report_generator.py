import json
import argparse
import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def generate_excel_report(input_json_path, output_excel_path):
    print(f"[Agent 3] Reading trace data from: {input_json_path}")
    
    if not os.path.exists(input_json_path):
        print(f"[Agent 3] ERROR: Input file not found: {input_json_path}")
        return False

    with open(input_json_path, 'r', encoding='utf-8') as f:
        trace_data = json.load(f)

    # Setup workbook and sheet
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Traceability Matrix"
    ws.views.sheetView[0].showGridLines = True

    # Styles
    title_font = Font(name="Segoe UI", size=16, bold=True, color="FFFFFF")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Segoe UI", size=10)
    
    blue_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    light_blue_fill = PatternFill(start_color="DCE6F1", end_color="DCE6F1", fill_type="solid")
    green_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
    red_fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")

    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )

    # Title Block
    ws.merge_cells("A1:E1")
    ws["A1"] = "SYSTEM REQUIREMENT TO TEST CASE TRACEABILITY REPORT"
    ws["A1"].font = title_font
    ws["A1"].fill = blue_fill
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions.height = 40

    # Summary Row
    ws["A3"] = "Export Date:"
    ws["A3"].font = Font(name="Segoe UI", size=10, bold=True)
    ws["B3"] = "2026-09-18 (Friday)"
    ws["B3"].font = data_font

    # Headers
    headers = ["Requirement ID", "Requirement Title", "Status/Type", "Mapped Test Cases", "Coverage Status"]
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=5, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = blue_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border
    ws.row_dimensions.height = 25

    # Data Rows
    current_row = 6
    total_reqs = 0
    covered_reqs = 0

    # Parse trace items (handle either flat list or nested dict formats)
    items = trace_data if isinstance(trace_data, list) else trace_data.get("traceability_matrix", trace_data.get("matrix", []))

    for item in items:
        req_id = item.get("id") or item.get("requirement_id", "N/A")
        title = item.get("title") or item.get("requirement_title", "N/A")
        status = item.get("status") or item.get("type", "Active")
        
        # Extract mapped test cases (support string list, or list of dicts with 'testcase_id' or 'id')
        test_cases = item.get("test_cases") or item.get("mapped_test_cases") or item.get("linked_testcases", [])
        if isinstance(test_cases, list):
            extracted_tcs = []
            for tc in test_cases:
                if isinstance(tc, dict):
                    tc_id = tc.get("testcase_id") or tc.get("id") or "Unknown TC"
                else:
                    tc_id = str(tc)
                if tc_id and tc_id not in ["N/A", "Unknown TC", "No test cases mapped"]:
                    extracted_tcs.append(tc_id)
            test_cases_str = ", ".join(extracted_tcs) if extracted_tcs else "No test cases mapped"
            is_covered = len(extracted_tcs) > 0
        else:
            test_cases_str = str(test_cases) if test_cases else "No test cases mapped"
            is_covered = bool(test_cases_str and test_cases_str not in ["No test cases mapped", "N/A"])
            
        coverage_status = "Covered" if is_covered else "GAP (Uncovered)"
        
        total_reqs += 1
        if is_covered:
            covered_reqs += 1

        row_data = [req_id, title, status, test_cases_str if test_cases_str else "No test cases mapped", coverage_status]
        
        for col_idx, value in enumerate(row_data, 1):
            cell = ws.cell(row=current_row, column=col_idx, value=value)
            cell.font = data_font
            cell.border = thin_border
            
            # Align IDs and Status
            if col_idx in [1, 3]:  # Align IDs and Status columns to center
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

            # Format the Coverage column
            if col_idx == 5:
                cell.fill = green_fill if is_covered else red_fill
                cell.font = Font(name="Segoe UI", size=10, bold=True, color="375623" if is_covered else "C00000")

        current_row += 1

    # Auto-adjust Column Widths
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.row > 1 and cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    # Save Excel file
    os.makedirs(os.path.dirname(output_excel_path), exist_ok=True)
    try:
        wb.save(output_excel_path)
        print(f"[Agent 3] Excel report generated successfully at: {output_excel_path}")
    except PermissionError:
        fallback_path = output_excel_path.replace(".xlsx", "_fallback.xlsx")
        print(f"[Agent 3] WARNING: Permission denied on {output_excel_path} (likely open in Excel).")
        print(f"[Agent 3] Saving to fallback location: {fallback_path}")
        wb.save(fallback_path)
        print(f"[Agent 3] Excel report generated successfully at: {fallback_path}")
    cov_pct = (covered_reqs / total_reqs) * 100 if total_reqs > 0 else 0
    print(f"[Agent 3] Coverage Summary: {covered_reqs}/{total_reqs} ({cov_pct:.2f}% Coverage)")
    return True

if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Convert traceability matrix JSON into Excel Report.")
    parser.add_argument("--input-json", default=r"C:\Users\trm1cob\.gemini\tmp\trm1cob\traceability_matrix.json")
    parser.add_argument("--output-excel", default=r"C:\Users\trm1cob\.gemini\tmp\trm1cob\traceability_report.xlsx")
    args = parser.parse_args()
    
    generate_excel_report(args.input_json, args.output_excel)
