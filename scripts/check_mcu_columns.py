import os
import sys
import openpyxl
import config

if len(sys.argv) < 2:
    print(f"[-] Error: Missing required file path argument.")
    print(f"[*] Usage: python {os.path.basename(sys.argv[0])} <path_to_excel_file.xlsx>")
    sys.exit(1)

path = sys.argv[1]

if not os.path.exists(path):
    print(f"[-] Error: File not found: {path}")
    sys.exit(1)
wb = openpyxl.load_workbook(path, data_only=True)
sheet = wb.active

headers = [str(cell).strip() if cell is not None else f"Col{idx}" for idx, cell in enumerate(next(sheet.iter_rows(values_only=True)), start=1)]

is_heading_col = headers.index("isHeading") if "isHeading" in headers else -1
artifact_type_col = headers.index("Artifact Type") if "Artifact Type" in headers else -1
section_col = headers.index("section") if "section" in headers else -1
id_col = headers.index("id") if "id" in headers else -1
primary_text_col = headers.index("Primary Text") if "Primary Text" in headers else -1

print(f"isHeading col: {is_heading_col}, Artifact Type col: {artifact_type_col}, section col: {section_col}")

heading_values = set()
artifact_types = set()
section_values = set()

reqs_by_section = {}

for r_idx, row in enumerate(sheet.iter_rows(values_only=True), start=1):
    if r_idx == 1:
        continue
    is_heading_val = row[is_heading_col] if is_heading_col != -1 else None
    art_type_val = row[artifact_type_col] if artifact_type_col != -1 else None
    sec_val = row[section_col] if section_col != -1 else None
    
    heading_values.add(is_heading_val)
    artifact_types.add(art_type_val)
    section_values.add(sec_val)
    
    if sec_val:
        reqs_by_section[sec_val] = reqs_by_section.get(sec_val, 0) + 1

print(f"Unique isHeading values: {heading_values}")
print(f"Unique Artifact Type values: {artifact_types}")
print(f"Unique Section values: {section_values}")
print(f"Requirements count by section:")
for s, count in sorted(reqs_by_section.items()):
    print(f"  - '{s}': {count}")
